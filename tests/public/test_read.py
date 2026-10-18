"""`graphme.read`, once a format: an invented file made here, its pages drawn, a stand-in model.

The model answers every question with nothing found, so this proves the reading runs end to
end -- rendered, asked, the file's words kept -- never what a real model would make of it."""
from __future__ import annotations

import json
import os
from pathlib import Path
from typing import Any

import httpx
import pytest

import graphme
from graphme import api, model_endpoint

WORDS = ("Invented heading", "Invented body text.")


def _settings(tmp_path: Path) -> graphme.Settings:
    chat = model_endpoint.PROVIDERS[model_endpoint.DEFAULT_PROVIDER]
    renderer = {name: value for name, value in os.environ.items() if name in (
        "GRAPHME_VIEWER_DIR", "GRAPHME_CHROME_EXECUTABLE")}
    return graphme.load_settings({
        **renderer, "GRAPHME_PROFILE": "test", "GRAPHME_EMBEDDING_PROVIDER": "hash",
        "GRAPHME_DATA_DIR": str(tmp_path / "data"), chat.key_env: "invented-key"})


def _write(kind: str, path: Path) -> None:
    if kind == "pdf":
        import pymupdf

        pdf = pymupdf.open()
        page = pdf.new_page()
        for line, text in enumerate(WORDS):
            page.insert_text((72, 72 + 28 * line), text)
        pdf.save(path)
    elif kind == "docx":
        from docx import Document

        document = Document()
        for text in WORDS:
            document.add_paragraph(text)
        document.save(str(path))
    elif kind == "xlsx":
        from openpyxl import Workbook

        workbook = Workbook()
        for row, text in enumerate(WORDS, start=1):
            workbook.active.cell(row=row, column=1, value=text)
        workbook.save(path)
    else:
        from pptx import Presentation
        from pptx.util import Inches

        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        for line, text in enumerate(WORDS):
            box = slide.shapes.add_textbox(Inches(1), Inches(1 + line), Inches(5), Inches(1))
            box.text_frame.text = text
        deck.save(str(path))


@pytest.mark.parametrize("kind", ["pdf", "docx", "xlsx", "pptx"])
def test_a_document_is_read_with_its_words_and_its_pages_shown_to_the_model(
    kind: str, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    settings = _settings(tmp_path)
    if kind != "pdf" and not (
        settings.chrome_executable
        and (settings.viewer_directory / "dist").is_dir()
    ):
        pytest.skip("DOCX, XLSX and PPTX are drawn by Chrome with the viewer bundle (README 0-1)")
    asked: list[dict[str, Any]] = []

    def nothing_found(request: httpx.Request) -> httpx.Response:
        body = json.loads(request.content)
        asked.append(body)
        schema = body["response_format"]["json_schema"]["schema"]
        answer = json.dumps({field: [] for field in schema["required"]})
        return httpx.Response(200, json={"choices": [{"message": {"content": answer}}]})

    client = httpx.Client
    monkeypatch.setattr(api.httpx, "Client", lambda **_: client(
        transport=httpx.MockTransport(nothing_found)))
    source = tmp_path / f"invented.{kind}"
    _write(kind, source)

    read = graphme.read(source, settings=settings)

    assert read.document.file_name == source.name
    assert [node.label for node in read.nodes if node.type.value == "text_chunk"] == list(WORDS)
    shown = [part for body in asked for message in body["messages"]
             if isinstance(message["content"], list) for part in message["content"]
             if part["type"] == "image_url"]
    assert shown, "the model is shown the drawn page"
