"""A chart in a Word file is a chart in the graph.

Word keeps a chart as a part of its own, drawn by the layout like a picture but holding no
picture: read for pictures alone, it never became a node, had no crop and no caption, and
the outline went on as if the page held nothing there.
"""
from __future__ import annotations

from pathlib import Path

import docx
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.oxml import parse_xml
from pptx.chart.data import CategoryChartData
from pptx.chart.xmlwriter import ChartXmlWriter
from pptx.enum.chart import XL_CHART_TYPE

from graphme.adapters.extractors.docx import DocxExtractor
from graphme.adapters.rendering.browser_office import RenderedPage
from graphme.application.page_units import flow_pages
from graphme.domain.enums import DocumentFormat, NodeType

CHART_TYPE = "application/vnd.openxmlformats-officedocument.drawingml.chart+xml"
INLINE = (
    '<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
    'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
    'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
    'xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart" '
    'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
    '<w:drawing><wp:inline><wp:extent cx="4000000" cy="2500000"/><wp:docPr id="1" name="Chart 1"/>'
    '<a:graphic><a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/chart">'
    '<c:chart r:id="{rid}"/></a:graphicData></a:graphic></wp:inline></w:drawing></w:r>'
)


def _memo_with_a_chart(path: Path) -> Path:
    data = CategoryChartData()
    data.categories = ["Q1", "Q2", "Q3"]
    data.add_series("Returns", (12, 9, 4))
    memo = docx.Document()
    memo.add_paragraph("Returns fell each quarter after the seal was changed.")
    holder = memo.add_paragraph()
    blob = ChartXmlWriter(XL_CHART_TYPE.COLUMN_CLUSTERED, data).xml.encode()
    part = Part(PackURI("/word/charts/chart1.xml"), CHART_TYPE, blob, memo.part.package)
    holder._p.append(parse_xml(INLINE.format(rid=memo.part.relate_to(part, RT.CHART))))
    memo.save(path)
    return path


def test_a_chart_in_a_word_file_becomes_a_chart_node_with_its_values(tmp_path: Path) -> None:
    extracted = DocxExtractor().extract(_memo_with_a_chart(tmp_path / "memo.docx"))

    charts = [n for n in extracted.nodes if n.type is NodeType.CHART]
    assert len(charts) == 1
    chart = charts[0]
    assert chart.metadata["body_index"] == 2 and chart.metadata["chart_index"] == 1
    assert chart.metadata["categories"] == ["Q1", "Q2", "Q3"]
    assert chart.metadata["series"] == [{"name": "Returns", "values": [12.0, 9.0, 4.0]}]


def test_a_chart_is_placed_on_the_page_the_layout_drew_it(tmp_path: Path) -> None:
    extracted = DocxExtractor().extract(_memo_with_a_chart(tmp_path / "memo.docx"))
    chart = next(n for n in extracted.nodes if n.type is NodeType.CHART)
    # The renderer counts body elements from zero; the chart hangs from the second.
    page = RenderedPage(index=2, width=600, height=800, png=b"", elements=(
        {"kind": "chart", "story": "body", "instance": "body", "path": [1, 0],
         "x": 60, "y": 100, "w": 400, "h": 250},
    ))

    assert flow_pages(extracted, [page], DocumentFormat.DOCX).get(chart.id) == 2
