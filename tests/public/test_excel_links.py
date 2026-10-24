"""What a workbook's formulas and pictures say about each other. Invented workbooks only."""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.worksheet.table import Table

from graphme.adapters.extractors.excel import ExcelExtractor
from graphme.domain.enums import EdgeType


def test_the_same_cell_on_two_sheets_links_to_the_table_twice(tmp_path: Path) -> None:
    # One!A1 and Two!A1 both read the table; the second link was dropped as a repeat of the
    # first because the repeat check left out the sheet.
    book = Workbook()
    data = book.active
    data.title = "Data"
    for row in (["Item", "Amount"], ["Bolt", 3], ["Nut", 5]):
        data.append(row)
    data.add_table(Table(displayName="PartsTbl", ref="A1:B3"))
    for name in ("One", "Two"):
        book.create_sheet(name)["A1"] = "=SUM(Data!B2:B3)"
    book.save(tmp_path / "parts.xlsx")

    extracted = ExcelExtractor().extract(tmp_path / "parts.xlsx")

    uses = [e for e in extracted.edges if e.type is EdgeType.USES_VALUE_FROM]
    assert sorted(e.locator.sheet for e in uses if e.locator) == ["One", "Two"]


def test_a_jpg_picture_is_named_with_the_registered_media_type() -> None:
    # "image/jpg" is not a media type; JPEG's is "image/jpeg", whatever the file suffix.
    assert ExcelExtractor()._image_mime_type(".jpg") == "image/jpeg"
    assert ExcelExtractor()._image_mime_type(".JPEG") == "image/jpeg"


def test_a_formula_reading_a_sheet_named_in_korean_links_to_its_table(tmp_path: Path) -> None:
    # Excel writes a sheet name of letters bare, whatever the script: =SUM(재고!B2:B3). A
    # reference read as Latin letters only found no sheet there, and the link was lost.
    book = Workbook()
    data = book.active
    data.title = "재고"
    for row in (["품목", "수량"], ["볼트", 3], ["너트", 5]):
        data.append(row)
    data.add_table(Table(displayName="Stock", ref="A1:B3"))
    book.create_sheet("합계")["A1"] = "=SUM(재고!B2:B3)"
    book.save(tmp_path / "stock.xlsx")

    extracted = ExcelExtractor().extract(tmp_path / "stock.xlsx")

    uses = [e for e in extracted.edges if e.type is EdgeType.USES_VALUE_FROM]
    assert [e.locator.sheet for e in uses if e.locator] == ["합계"]
