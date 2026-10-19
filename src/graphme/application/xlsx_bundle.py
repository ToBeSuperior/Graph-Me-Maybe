from __future__ import annotations

import html
import json
import math
import os
import posixpath
import shutil
import tempfile
import xml.etree.ElementTree as ET
import zipfile
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Final, Literal, cast

from openpyxl.utils import get_column_letter, range_boundaries
from openpyxl.worksheet.worksheet import Worksheet
from PIL import Image, ImageDraw

from graphme.adapters.extractors.excel import ExcelExtractor, open_workbook
from graphme.adapters.extractors.xlsx_drawings import (
    DrawingObject,
    parse_xlsx_drawings,
)
from graphme.adapters.offline import PixelSegmenter
from graphme.adapters.rendering.browser_office import BrowserOfficeRenderer, SourcePageMapping
from graphme.adapters.rendering.page_image import PageImage, store_pages
from graphme.adapters.storage.assets import LocalAssetStore
from graphme.domain.enums import DocumentFormat, EdgeType
from graphme.domain.ids import canonical_json, sha256_file
from graphme.domain.models import ExtractedDocument, PageSegmentationRequest, SourceLocator

MappingStatus = Literal["mapped", "partial", "unmapped"]

_MAIN_NS: Final = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
_DOC_REL_NS: Final = (
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
)
_PKG_REL_NS: Final = "http://schemas.openxmlformats.org/package/2006/relationships"
_WORKSHEET_REL_TYPE: Final = (
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet"
)

ET.register_namespace("", _MAIN_NS)
ET.register_namespace("r", _DOC_REL_NS)
ET.register_namespace("mc", "http://schemas.openxmlformats.org/markup-compatibility/2006")
ET.register_namespace("x14ac", "http://schemas.microsoft.com/office/spreadsheetml/2009/9/ac")
ET.register_namespace("x15", "http://schemas.microsoft.com/office/spreadsheetml/2010/11/main")
ET.register_namespace("x15ac", "http://schemas.microsoft.com/office/spreadsheetml/2010/11/ac")


@dataclass(frozen=True, slots=True)
class PreparedSheet:
    name: str
    print_area: str | None
    mapping_status: MappingStatus
    page_start: int | None
    page_end: int | None
    detail: str | None = None


@dataclass(frozen=True, slots=True)
class WorkbookBundleResult:
    source: Path
    prepared_source: Path
    output_directory: Path
    mapping_status: MappingStatus
    sheets: tuple[PreparedSheet, ...]
    page_count: int
    node_count: int
    manifest: Path
    html: Path
    markdown: Path


class XlsxBundleOrchestrator:
    """Build a local, inspectable XLSX rendering and extraction bundle."""

    def __init__(
        self,
        *,
        browser_renderer: BrowserOfficeRenderer | None = None,
        tile_size: int = 1024,
    ) -> None:
        if tile_size <= 0:
            raise ValueError("tile size must be positive")
        self._browser_renderer = browser_renderer
        self._tile_size = tile_size

    def build(self, source: Path, output_directory: Path) -> WorkbookBundleResult:
        source = source.resolve()
        if not source.is_file():
            raise FileNotFoundError(source)
        if source.suffix.lower() != ".xlsx":
            raise ValueError("XLSX bundle input must use the .xlsx extension")

        output_directory = output_directory.resolve()
        output_directory.mkdir(parents=True, exist_ok=True)
        prepared_directory = output_directory / "prepared"
        page_directory = output_directory / "pages"
        tile_directory = output_directory / "tiles"
        overlay_directory = output_directory / "overlays"
        asset_directory = output_directory / "assets"
        for directory in (
            prepared_directory,
            page_directory,
            tile_directory,
            overlay_directory,
            asset_directory,
        ):
            directory.mkdir(parents=True, exist_ok=True)

        source_copy = output_directory / source.name
        prepared_source = prepared_directory / f"{source.stem}.prepared.xlsx"
        shutil.copy2(source, source_copy)
        shutil.copy2(source, prepared_source)
        sheet_areas = self._prepare_copy(prepared_source)

        asset_store = LocalAssetStore(asset_directory)
        if self._browser_renderer is None:
            raise RuntimeError("a browser renderer is required to render a workbook")
        render_result = store_pages(
            self._browser_renderer.render_with_mapping(prepared_source, DocumentFormat.XLSX),
            asset_store,
        )
        pages = render_result.pages
        mappings = render_result.mappings
        mapping_by_sheet = {
            mapping.locator.sheet: mapping
            for mapping in mappings
            if mapping.locator.sheet is not None
        }
        sheets = self._sheet_statuses(sheet_areas, mapping_by_sheet)
        mapping_status = self._whole_status(sheets)

        page_paths = self._materialize_pages(pages, asset_store, page_directory)
        segmenter = PixelSegmenter()
        segmentation: dict[int, list[dict[str, object]]] = {}
        tiles: dict[int, list[dict[str, object]]] = {}
        overlays: dict[int, str] = {}
        for page, page_path in zip(pages, page_paths, strict=True):
            payload = page_path.read_bytes()
            locator = self._page_locator(prepared_source.name, page.page_index, mappings)
            request = PageSegmentationRequest(
                page_index=page.page_index,
                width=page.width,
                height=page.height,
                asset_hash=page.asset_hash,
                locator=locator,
            )
            result = segmenter.segment_page(request=request, image=payload)
            units = [unit.model_dump(mode="json") for unit in result.units]
            segmentation[page.page_index] = cast(list[dict[str, object]], units)
            tiles[page.page_index] = self._make_tiles(page_path, tile_directory, page.page_index)
            overlays[page.page_index] = self._make_overlay(
                page_path, overlay_directory, page.page_index, units
            )

        extracted = ExcelExtractor(asset_store=asset_store).extract(prepared_source)
        hierarchy = self._hierarchy_data(extracted)
        hierarchy_path = output_directory / "hierarchy.json"
        hierarchy_path.write_text(
            canonical_json(hierarchy) + "\n", encoding="utf-8"
        )

        manifest_data = {
            "schema_version": 1,
            "source": {
                "original_path": str(source),
                "copied_path": source_copy.name,
                "prepared_path": str(prepared_source.relative_to(output_directory)),
                "sha256": sha256_file(source),
            },
            "preparation": {
                "landscape": True,
                "fit_to_width": 1,
                "fit_to_height": 0,
                "sheets": [asdict(sheet) for sheet in sheets],
            },
            "mapping": {
                "status": mapping_status,
                "entries": [self._mapping_data(mapping) for mapping in mappings],
            },
            "pages": [
                {
                    **asdict(page),
                    "path": str(path.relative_to(output_directory)),
                    "overlay": overlays[page.page_index],
                    "tiles": tiles[page.page_index],
                    "offline_regions": segmentation[page.page_index],
                }
                for page, path in zip(pages, page_paths, strict=True)
            ],
            "hierarchy": {
                "path": hierarchy_path.name,
                "nodes": len(extracted.nodes),
                "edges": len(extracted.edges),
                "evidence": len(extracted.evidence),
                "warnings": len(extracted.warnings),
            },
            "providers": {
                "renderer": "ooxml-browser",
                "external_vlm_called": False,
                "layout": "deterministic-local",
            },
        }
        manifest_path = output_directory / "manifest.json"
        manifest_path.write_text(
            json.dumps(manifest_data, ensure_ascii=False, indent=2) + "\n",
            encoding="utf-8",
        )
        markdown_path = output_directory / "README.md"
        html_path = output_directory / "index.html"
        markdown_path.write_text(
            self._markdown(source.name, mapping_status, sheets, pages, extracted),
            encoding="utf-8",
        )
        html_path.write_text(
            self._html(source.name, mapping_status, sheets, pages, segmentation, hierarchy),
            encoding="utf-8",
        )
        return WorkbookBundleResult(
            source=source,
            prepared_source=prepared_source,
            output_directory=output_directory,
            mapping_status=mapping_status,
            sheets=sheets,
            page_count=len(pages),
            node_count=len(extracted.nodes),
            manifest=manifest_path,
            html=html_path,
            markdown=markdown_path,
        )

    @classmethod
    def _prepare_copy(cls, prepared_source: Path) -> dict[str, str | None]:
        workbook = open_workbook(prepared_source, data_only=False, read_only=False)
        drawing_manifest = parse_xlsx_drawings(prepared_source)
        by_sheet: dict[str, list[DrawingObject]] = {}
        for drawing in drawing_manifest.objects:
            by_sheet.setdefault(drawing.worksheet, []).append(drawing)
        try:
            areas = {
                sheet.title: cls._used_and_drawing_range(
                    sheet, by_sheet.get(sheet.title, [])
                )
                for sheet in workbook.worksheets
            }
        finally:
            workbook.close()

        cls._patch_print_settings(prepared_source, areas)
        return areas

    @classmethod
    def _patch_print_settings(
        cls, prepared_source: Path, areas: dict[str, str | None]
    ) -> None:
        with zipfile.ZipFile(prepared_source, "r") as archive:
            members = archive.infolist()
            names = [member.filename for member in members]
            if len(names) != len(set(names)):
                raise ValueError("XLSX package contains duplicate member names")
            workbook_xml = archive.read("xl/workbook.xml")
            relationships_xml = archive.read("xl/_rels/workbook.xml.rels")
            workbook_root = ET.fromstring(workbook_xml)
            relationships_root = ET.fromstring(relationships_xml)
            sheet_parts = cls._worksheet_parts(workbook_root, relationships_root)
            sheet_names = [name for name, _ in sheet_parts]
            if sheet_names != list(areas):
                raise ValueError("Workbook worksheet order changed during preparation")

            replacements = {
                "xl/workbook.xml": cls._patch_workbook_print_areas(
                    workbook_root, areas
                )
            }
            for _sheet_name, part_name in sheet_parts:
                sheet_root = ET.fromstring(archive.read(part_name))
                cls._patch_worksheet_page_setup(sheet_root)
                replacements[part_name] = ET.tostring(
                    sheet_root, encoding="utf-8", xml_declaration=True
                )

            file_descriptor, temporary_name = tempfile.mkstemp(
                prefix=f".{prepared_source.name}.",
                suffix=".tmp",
                dir=prepared_source.parent,
            )
            os.close(file_descriptor)
            temporary_path = Path(temporary_name)
            try:
                with zipfile.ZipFile(temporary_path, "w") as output:
                    output.comment = archive.comment
                    for member in members:
                        output.writestr(
                            member,
                            replacements.get(member.filename, archive.read(member)),
                        )
            except BaseException:
                temporary_path.unlink(missing_ok=True)
                raise
        # Replaced once the source is closed: Windows refuses to replace a file still open.
        try:
            os.replace(temporary_path, prepared_source)
        finally:
            temporary_path.unlink(missing_ok=True)

    @staticmethod
    def _worksheet_parts(
        workbook_root: ET.Element, relationships_root: ET.Element
    ) -> list[tuple[str, str]]:
        relationships = {
            relationship.attrib["Id"]: relationship
            for relationship in relationships_root.findall(f"{{{_PKG_REL_NS}}}Relationship")
        }
        result: list[tuple[str, str]] = []
        sheets = workbook_root.find(f"{{{_MAIN_NS}}}sheets")
        if sheets is None:
            raise ValueError("XLSX workbook has no worksheets")
        for sheet in sheets.findall(f"{{{_MAIN_NS}}}sheet"):
            relationship_id = sheet.get(f"{{{_DOC_REL_NS}}}id")
            relationship = relationships.get(relationship_id or "")
            if relationship is None or relationship.get("Type") != _WORKSHEET_REL_TYPE:
                raise ValueError(
                    f"Unresolved worksheet relationship for {sheet.get('name', '')}"
                )
            target = relationship.get("Target")
            if not target or "?" in target or "#" in target:
                raise ValueError("Unsafe worksheet relationship target")
            part_name = (
                posixpath.normpath(target).lstrip("/")
                if target.startswith("/")
                else posixpath.normpath(posixpath.join("xl", target))
            )
            if part_name.startswith("../") or not part_name.startswith("xl/"):
                raise ValueError("Worksheet relationship escapes the XLSX package")
            result.append((sheet.get("name", ""), part_name))
        return result

    @staticmethod
    def _patch_workbook_print_areas(
        workbook_root: ET.Element, areas: dict[str, str | None]
    ) -> bytes:
        defined_names = workbook_root.find(f"{{{_MAIN_NS}}}definedNames")
        if defined_names is None:
            defined_names = ET.Element(f"{{{_MAIN_NS}}}definedNames")
            children = list(workbook_root)
            insertion_index = next(
                (
                    index
                    for index, child in enumerate(children)
                    if child.tag
                    in {
                        f"{{{_MAIN_NS}}}calcPr",
                        f"{{{_MAIN_NS}}}oleSize",
                        f"{{{_MAIN_NS}}}customWorkbookViews",
                        f"{{{_MAIN_NS}}}pivotCaches",
                        f"{{{_MAIN_NS}}}smartTagPr",
                        f"{{{_MAIN_NS}}}smartTagTypes",
                        f"{{{_MAIN_NS}}}webPublishing",
                        f"{{{_MAIN_NS}}}fileRecoveryPr",
                        f"{{{_MAIN_NS}}}webPublishObjects",
                        f"{{{_MAIN_NS}}}extLst",
                    }
                ),
                len(children),
            )
            workbook_root.insert(insertion_index, defined_names)
        for defined_name in list(defined_names):
            if defined_name.get("name") == "_xlnm.Print_Area":
                defined_names.remove(defined_name)
        for local_sheet_id, (sheet_name, area) in enumerate(areas.items()):
            if area is None:
                continue
            defined_name = ET.SubElement(
                defined_names,
                f"{{{_MAIN_NS}}}definedName",
                {"name": "_xlnm.Print_Area", "localSheetId": str(local_sheet_id)},
            )
            defined_name.text = (
                f"'{sheet_name.replace(chr(39), chr(39) * 2)}'!"
                f"{XlsxBundleOrchestrator._absolute_range(area)}"
            )
        if not list(defined_names):
            workbook_root.remove(defined_names)
        return cast(
            bytes,
            ET.tostring(workbook_root, encoding="utf-8", xml_declaration=True),
        )

    @staticmethod
    def _absolute_range(area: str) -> str:
        min_col, min_row, max_col, max_row = range_boundaries(area)
        start = f"${get_column_letter(min_col)}${min_row}"
        end = f"${get_column_letter(max_col)}${max_row}"
        return start if start == end else f"{start}:{end}"

    @staticmethod
    def _patch_worksheet_page_setup(sheet_root: ET.Element) -> None:
        sheet_properties = sheet_root.find(f"{{{_MAIN_NS}}}sheetPr")
        if sheet_properties is None:
            sheet_properties = ET.Element(f"{{{_MAIN_NS}}}sheetPr")
            sheet_root.insert(0, sheet_properties)
        setup_properties = sheet_properties.find(f"{{{_MAIN_NS}}}pageSetUpPr")
        if setup_properties is None:
            setup_properties = ET.SubElement(
                sheet_properties, f"{{{_MAIN_NS}}}pageSetUpPr"
            )
        setup_properties.set("fitToPage", "1")

        page_setup = sheet_root.find(f"{{{_MAIN_NS}}}pageSetup")
        if page_setup is None:
            page_setup = ET.Element(f"{{{_MAIN_NS}}}pageSetup")
            children = list(sheet_root)
            insertion_index = next(
                (
                    index
                    for index, child in enumerate(children)
                    if child.tag
                    in {
                        f"{{{_MAIN_NS}}}headerFooter",
                        f"{{{_MAIN_NS}}}rowBreaks",
                        f"{{{_MAIN_NS}}}colBreaks",
                        f"{{{_MAIN_NS}}}customProperties",
                        f"{{{_MAIN_NS}}}cellWatches",
                        f"{{{_MAIN_NS}}}ignoredErrors",
                        f"{{{_MAIN_NS}}}smartTags",
                        f"{{{_MAIN_NS}}}drawing",
                        f"{{{_MAIN_NS}}}legacyDrawing",
                        f"{{{_MAIN_NS}}}legacyDrawingHF",
                        f"{{{_MAIN_NS}}}picture",
                        f"{{{_MAIN_NS}}}oleObjects",
                        f"{{{_MAIN_NS}}}controls",
                        f"{{{_MAIN_NS}}}webPublishItems",
                        f"{{{_MAIN_NS}}}tableParts",
                        f"{{{_MAIN_NS}}}extLst",
                    }
                ),
                len(children),
            )
            sheet_root.insert(insertion_index, page_setup)
        page_setup.attrib.pop("scale", None)
        page_setup.set("orientation", "landscape")
        page_setup.set("fitToWidth", "1")
        page_setup.set("fitToHeight", "0")

    @staticmethod
    def _used_and_drawing_range(
        sheet: Worksheet, drawings: list[DrawingObject]
    ) -> str | None:
        bounds: list[tuple[int, int, int, int]] = []
        for cell in sheet._cells.values():
            if cell.value is not None:
                bounds.append((cell.column, cell.row, cell.column, cell.row))
        for table in sheet.tables.values():
            bounds.append(cast(tuple[int, int, int, int], range_boundaries(table.ref)))
        for merged in sheet.merged_cells.ranges:
            bounds.append((merged.min_col, merged.min_row, merged.max_col, merged.max_row))
        for drawing in drawings:
            start = drawing.anchor_from
            if start is None or start.col is None or start.row is None:
                continue
            min_col, min_row = start.col + 1, start.row + 1
            end = drawing.anchor_to
            if end is not None and end.col is not None and end.row is not None:
                max_col, max_row = end.col + 1, end.row + 1
            else:
                max_col, max_row = min_col, min_row
            bounds.append((min_col, min_row, max_col, max_row))
        if not bounds:
            return None
        min_col = min(bound[0] for bound in bounds)
        min_row = min(bound[1] for bound in bounds)
        max_col = max(bound[2] for bound in bounds)
        max_row = max(bound[3] for bound in bounds)
        start_cell = f"{get_column_letter(min_col)}{min_row}"
        end_cell = f"{get_column_letter(max_col)}{max_row}"
        return start_cell if start_cell == end_cell else f"{start_cell}:{end_cell}"

    @staticmethod
    def _sheet_statuses(
        areas: dict[str, str | None],
        mapping_by_sheet: dict[str, SourcePageMapping],
    ) -> tuple[PreparedSheet, ...]:
        result = []
        for name, area in areas.items():
            mapping = mapping_by_sheet.get(name)
            if area is None:
                status: MappingStatus = "unmapped"
                detail = "blank worksheet excluded from rendering"
            elif mapping is None:
                status = "unmapped"
                detail = "no worksheet page mapping"
            else:
                status = "mapped"
                detail = None
            result.append(
                PreparedSheet(
                    name=name,
                    print_area=area,
                    mapping_status=status,
                    page_start=mapping.page_start if mapping else None,
                    page_end=mapping.page_end if mapping else None,
                    detail=detail,
                )
            )
        return tuple(result)

    @staticmethod
    def _whole_status(sheets: tuple[PreparedSheet, ...]) -> MappingStatus:
        meaningful = [sheet for sheet in sheets if sheet.print_area is not None]
        mapped = sum(sheet.mapping_status == "mapped" for sheet in meaningful)
        if meaningful and mapped == len(meaningful):
            return "mapped"
        if mapped:
            return "partial"
        return "unmapped"

    @staticmethod
    def _materialize_pages(
        pages: tuple[PageImage, ...], store: LocalAssetStore, directory: Path
    ) -> list[Path]:
        paths = []
        for page in pages:
            target = directory / f"page-{page.page_index:04d}.png"
            target.write_bytes(store.read(page.asset_hash, ".png"))
            paths.append(target)
        return paths

    def _make_tiles(
        self, page_path: Path, directory: Path, page_index: int
    ) -> list[dict[str, object]]:
        page_dir = directory / f"page-{page_index:04d}"
        page_dir.mkdir(parents=True, exist_ok=True)
        entries: list[dict[str, object]] = []
        with Image.open(page_path) as image:
            rgb = image.convert("RGB")
            columns = math.ceil(rgb.width / self._tile_size)
            rows = math.ceil(rgb.height / self._tile_size)
            for row in range(rows):
                for column in range(columns):
                    box = (
                        column * self._tile_size,
                        row * self._tile_size,
                        min((column + 1) * self._tile_size, rgb.width),
                        min((row + 1) * self._tile_size, rgb.height),
                    )
                    path = page_dir / f"tile-r{row:03d}-c{column:03d}.png"
                    rgb.crop(box).save(path, format="PNG")
                    entries.append(
                        {
                            "row": row,
                            "column": column,
                            "pixel_bbox": list(box),
                            "path": str(path.relative_to(directory.parent)),
                        }
                    )
        return entries

    @staticmethod
    def _make_overlay(
        page_path: Path,
        directory: Path,
        page_index: int,
        units: list[dict[str, object]],
    ) -> str:
        with Image.open(page_path) as image:
            base = image.convert("RGBA")
            annotations = Image.new("RGBA", base.size, (0, 0, 0, 0))
            draw = ImageDraw.Draw(annotations, "RGBA")
            for unit in units:
                bbox = cast(dict[str, float], unit["bbox"])
                box = (
                    round(bbox["x0"] * base.width),
                    round(bbox["y0"] * base.height),
                    round(bbox["x1"] * base.width),
                    round(bbox["y1"] * base.height),
                )
                draw.rectangle(box, outline=(220, 38, 38, 255), width=3)
                label_box = (box[0], box[1], box[0] + 24, box[1] + 18)
                draw.rectangle(label_box, fill=(255, 255, 255, 224))
                draw.text(
                    (box[0] + 4, box[1] + 2),
                    str(unit["reading_order"]),
                    fill=(153, 27, 27, 255),
                )
            target = directory / f"page-{page_index:04d}.png"
            Image.alpha_composite(base, annotations).convert("RGB").save(
                target, format="PNG"
            )
        return str(target.relative_to(directory.parent))

    @staticmethod
    def _page_locator(
        file_name: str,
        page_index: int,
        mappings: tuple[SourcePageMapping, ...],
    ) -> SourceLocator:
        mapping = next(
            (
                item
                for item in mappings
                if item.page_start <= page_index <= item.page_end
            ),
            None,
        )
        return SourceLocator(
            file_name=file_name,
            label=mapping.locator.label if mapping else f"Page {page_index}",
            page=page_index,
            sheet=mapping.locator.sheet if mapping else None,
            cell_range=mapping.locator.cell_range if mapping else None,
        )

    @staticmethod
    def _mapping_data(mapping: SourcePageMapping) -> dict[str, object]:
        return {
            "locator": mapping.locator.model_dump(mode="json"),
            "page_start": mapping.page_start,
            "page_end": mapping.page_end,
        }

    @staticmethod
    def _hierarchy_data(extracted: ExtractedDocument) -> dict[str, object]:
        children: dict[str, list[str]] = {}
        for edge in extracted.edges:
            if edge.type is EdgeType.CONTAINS:
                children.setdefault(edge.source_id, []).append(edge.target_id)
        return {
            "document": extracted.document.model_dump(mode="json"),
            "nodes": [node.model_dump(mode="json") for node in extracted.nodes],
            "edges": [edge.model_dump(mode="json") for edge in extracted.edges],
            "evidence": [item.model_dump(mode="json") for item in extracted.evidence],
            "warnings": [warning.model_dump(mode="json") for warning in extracted.warnings],
            "children": children,
        }

    @staticmethod
    def _markdown(
        name: str,
        status: MappingStatus,
        sheets: tuple[PreparedSheet, ...],
        pages: tuple[PageImage, ...],
        extracted: ExtractedDocument,
    ) -> str:
        rows = "\n".join(
            f"| {sheet.name} | {sheet.print_area or '—'} | {sheet.mapping_status} | "
            f"{sheet.page_start or '—'}–{sheet.page_end or '—'} |"
            for sheet in sheets
        )
        images = "\n".join(
            f"## Page {page.page_index}\n\n"
            f"![Page {page.page_index}](pages/page-{page.page_index:04d}.png)"
            for page in pages
        )
        return (
            f"# {name} workbook bundle\n\n"
            f"- Mapping status: **{status}**\n"
            f"- Rendered pages: **{len(pages)}**\n"
            f"- Native hierarchy nodes: **{len(extracted.nodes)}**\n"
            "- Layout extraction: deterministic offline pixels; no external VLM.\n\n"
            "| Worksheet | Print area | Mapping | Pages |\n"
            "|---|---|---|---|\n"
            f"{rows}\n\n{images}\n"
        )

    @staticmethod
    def _html(
        name: str,
        status: MappingStatus,
        sheets: tuple[PreparedSheet, ...],
        pages: tuple[PageImage, ...],
        segmentation: dict[int, list[dict[str, object]]],
        hierarchy: dict[str, object],
    ) -> str:
        sheet_rows = "".join(
            "<tr>"
            f"<td>{html.escape(sheet.name)}</td>"
            f"<td><code>{html.escape(sheet.print_area or '—')}</code></td>"
            f"<td><span class='status {sheet.mapping_status}'>{sheet.mapping_status}</span></td>"
            f"<td>{sheet.page_start or '—'}–{sheet.page_end or '—'}</td>"
            "</tr>"
            for sheet in sheets
        )
        page_cards = "".join(
            "<article class='page'>"
            f"<h2>Page {page.page_index}</h2>"
            f"<p>{len(segmentation[page.page_index])} offline region(s), "
            f"{page.width}×{page.height}px</p>"
            f"<img loading='lazy' src='overlays/page-{page.page_index:04d}.png' "
            f"alt='Page {page.page_index} with offline region overlay'>"
            "</article>"
            for page in pages
        )
        node_count = len(cast(list[object], hierarchy["nodes"]))
        return f"""<!doctype html>
<html lang="en">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<title>{html.escape(name)} workbook bundle</title>
<style>
:root {{ color-scheme: light dark; font-family: system-ui, sans-serif; }}
body {{ margin: 0; background: #f4f5f7; color: #17202a; }}
main {{ max-width: 1100px; margin: auto; padding: 2rem; }}
header {{ background: #17202a; color: white; padding: 2rem; border-radius: 1rem; }}
.summary {{ display: flex; gap: 1rem; flex-wrap: wrap; }}
.summary b {{ font-size: 1.5rem; }}
table {{ width: 100%; border-collapse: collapse; background: white; margin: 1.5rem 0; }}
th, td {{ padding: .7rem; border-bottom: 1px solid #d8dde3; text-align: left; }}
.status {{ padding: .2rem .5rem; border-radius: 999px; background: #e5e7eb; }}
.mapped {{ background: #d1fae5; color: #065f46; }}
.partial {{ background: #fef3c7; color: #92400e; }}
.unmapped {{ background: #fee2e2; color: #991b1b; }}
.page {{ background: white; padding: 1rem; margin: 1.5rem 0; border-radius: .75rem; }}
img {{ width: 100%; height: auto; }}
@media (prefers-color-scheme: dark) {{
  body {{ background: #111827; color: #e5e7eb; }}
  table, .page {{ background: #1f2937; }}
}}
</style>
</head>
<body><main>
<header>
<h1>{html.escape(name)}</h1>
<div class="summary">
<p>Mapping<br><b>{status}</b></p>
<p>Pages<br><b>{len(pages)}</b></p>
<p>Native nodes<br><b>{node_count}</b></p>
</div>
<p>Prepared and rendered locally. No external VLM was called.</p>
</header>
<table>
<thead><tr><th>Worksheet</th><th>Print area</th><th>Status</th><th>Pages</th></tr></thead>
<tbody>{sheet_rows}</tbody>
</table>
{page_cards}
</main></body>
</html>
"""
