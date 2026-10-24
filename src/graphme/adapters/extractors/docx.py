from __future__ import annotations

from collections.abc import Iterable, Iterator
from pathlib import Path
from typing import Any, cast

from docx import Document
from docx.document import Document as DocxDocument
from docx.oxml.ns import qn
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph
from pptx.chart.chart import Chart
from pptx.oxml import parse_xml as parse_chart

from graphme.adapters.extractors.powerpoint import chart_metadata
from graphme.domain.enums import DocumentFormat, EdgeType, EvidenceKind, NodeType
from graphme.domain.ids import canonical_locator, sha256_bytes, sha256_file, stable_id
from graphme.domain.models import (
    DocumentRecord,
    EvidenceRecord,
    ExtractedDocument,
    GraphEdge,
    GraphNode,
    JsonValue,
    SourceLocator,
    native_edge,
)
from graphme.ports.storage import AssetStore

CHART_TAG = "{http://schemas.openxmlformats.org/drawingml/2006/chart}chart"
"""Where a run draws a chart: the graphic's reference to the chart's own part."""


class DocxExtractor:
    """Extract body content from a DOCX using its native reading order."""

    format = DocumentFormat.DOCX

    def __init__(self, *, asset_store: AssetStore | None = None) -> None:
        self._asset_store = asset_store

    def extract(self, source: Path) -> ExtractedDocument:
        content_hash = sha256_file(source)
        document = Document(str(source))
        document_id = stable_id(
            content_hash, canonical_locator("document", {"file": source.name})
        )
        root_locator = SourceLocator(file_name=source.name, label=source.name)
        nodes: list[GraphNode] = [
            GraphNode(
                id=document_id,
                document_id=document_id,
                type=NodeType.DOCUMENT,
                label=source.name,
                locator=root_locator,
                content_hash=content_hash,
            )
        ]
        evidence: list[EvidenceRecord] = [
            EvidenceRecord(
                id=stable_id(content_hash, "evidence:document"),
                document_id=document_id,
                node_id=document_id,
                kind=EvidenceKind.SOURCE_METADATA,
                locator=root_locator,
                text=f"document {source.name}",
            )
        ]
        edges: list[GraphEdge] = []
        ordered_ids: list[str] = []

        for body_index, block in enumerate(
            _blocks(document.element.body, document._body), start=1
        ):
            if isinstance(block, Paragraph):
                self._extract_paragraph(
                    paragraph=block,
                    body_index=body_index,
                    document=document,
                    source=source,
                    content_hash=content_hash,
                    document_id=document_id,
                    nodes=nodes,
                    evidence=evidence,
                    ordered_ids=ordered_ids,
                )
            elif isinstance(block, Table):
                node, record = self._extract_table(
                    table=block,
                    body_index=body_index,
                    source=source,
                    content_hash=content_hash,
                    document_id=document_id,
                )
                nodes.append(node)
                evidence.append(record)
                ordered_ids.append(node.id)

        for child_id in ordered_ids:
            edges.append(
                native_edge(content_hash, document_id, document_id, child_id, EdgeType.CONTAINS)
            )

        extracted = ExtractedDocument(
            document=DocumentRecord(
                id=document_id,
                content_hash=content_hash,
                file_name=source.name,
                format=DocumentFormat.DOCX,
                source_path=source,
                size_bytes=source.stat().st_size,
                metadata={"native_extractor": "python-docx"},
            ),
            nodes=tuple(nodes),
            edges=tuple(edges),
            evidence=tuple(evidence),
        )
        return extracted.validate_canonical_semantic_graph()

    def _extract_paragraph(
        self,
        *,
        paragraph: Paragraph,
        body_index: int,
        document: DocxDocument,
        source: Path,
        content_hash: str,
        document_id: str,
        nodes: list[GraphNode],
        evidence: list[EvidenceRecord],
        ordered_ids: list[str],
    ) -> None:
        raw = _paragraph_text(paragraph)
        text = self._display_text(paragraph, raw)
        if text:
            # A paragraph is text whatever style it carries: reading "Heading 1" as a level
            # was the outline decided here, and an id that changed when someone restyled a
            # line was a second problem.
            node_id = stable_id(
                content_hash,
                canonical_locator("paragraph", {"body": body_index}),
            )
            locator = SourceLocator(
                file_name=source.name,
                label=f"body paragraph {body_index}",
                metadata={"body_index": body_index},
            )
            metadata: dict[str, JsonValue] = {
                "body_index": body_index,
                "style": paragraph.style.name if paragraph.style is not None else None,
                # A report indents with the paragraph or with spaces inside the line, and
                # uses both on one page. The display text is trimmed, so the second one
                # survives only as a count.
                "leading_spaces": len(raw) - len(raw.lstrip()),
            }
            nodes.append(
                GraphNode(
                    id=node_id,
                    document_id=document_id,
                    type=NodeType.TEXT_CHUNK,
                    label=text,
                    locator=locator,
                    text=text,
                    content_hash=sha256_bytes(text.encode()),
                    metadata=metadata,
                )
            )
            evidence.append(
                EvidenceRecord(
                    id=stable_id(content_hash, f"evidence:paragraph:{body_index}"),
                    document_id=document_id,
                    node_id=node_id,
                    kind=EvidenceKind.SOURCE_TEXT,
                    locator=locator,
                    text=text,
                )
            )
            ordered_ids.append(node_id)

        for box_index, box_text in enumerate(self._text_boxes(paragraph), start=1):
            # A floating text box is one unit the way a slide's text box is: the lines in
            # it belong together, and what encloses what is the hierarchy pass's question.
            box_id = stable_id(
                content_hash,
                canonical_locator("textbox", {"body": body_index, "box": box_index}),
            )
            box_locator = SourceLocator(
                file_name=source.name,
                label=f"text box {body_index}.{box_index}",
                metadata={
                    "story": "textbox",
                    "body_index": body_index,
                    "box_index": box_index,
                },
            )
            nodes.append(
                GraphNode(
                    id=box_id,
                    document_id=document_id,
                    type=NodeType.TEXT_CHUNK,
                    label=box_text.splitlines()[0][:200],
                    locator=box_locator,
                    text=box_text,
                    content_hash=sha256_bytes(box_text.encode()),
                    metadata={
                        "story": "textbox",
                        "body_index": body_index,
                        "box_index": box_index,
                    },
                )
            )
            evidence.append(
                EvidenceRecord(
                    id=stable_id(content_hash, f"evidence:textbox:{body_index}:{box_index}"),
                    document_id=document_id,
                    node_id=box_id,
                    kind=EvidenceKind.SOURCE_TEXT,
                    locator=box_locator,
                    text=box_text,
                )
            )
            ordered_ids.append(box_id)

        for image_index, (relationship_id, image_part) in enumerate(
            self._inline_images(paragraph, document), start=1
        ):
            image_id = stable_id(
                content_hash,
                canonical_locator(
                    "image", {"body": body_index, "image": image_index}
                ),
            )
            content_type = str(image_part.content_type)
            extension = self._extension(content_type)
            blob = bytes(image_part.blob)
            asset_hash = (
                self._asset_store.put(blob, suffix=extension)
                if self._asset_store is not None
                else None
            )
            locator = SourceLocator(
                file_name=source.name,
                label=f"inline image {body_index}.{image_index}",
                relationship_id=relationship_id,
                source_part=str(image_part.partname),
                metadata={"body_index": body_index, "image_index": image_index},
            )
            nodes.append(
                GraphNode(
                    id=image_id,
                    document_id=document_id,
                    type=NodeType.IMAGE,
                    label=f"Inline image {image_index}",
                    locator=locator,
                    content_hash=sha256_bytes(blob),
                    metadata={
                        "body_index": body_index,
                        "image_index": image_index,
                        "mime_type": content_type,
                        "size_bytes": len(blob),
                    },
                )
            )
            evidence.append(
                EvidenceRecord(
                    id=stable_id(content_hash, f"evidence:image:{body_index}:{image_index}"),
                    document_id=document_id,
                    node_id=image_id,
                    kind=EvidenceKind.SOURCE_MEDIA,
                    locator=locator,
                    data={"size_bytes": len(blob)} if asset_hash is None else None,
                    asset_hash=asset_hash,
                    mime_type=content_type,
                )
            )
            ordered_ids.append(image_id)

        for chart_index, (relationship_id, chart_part) in enumerate(
            self._charts(paragraph, document), start=1
        ):
            # A chart is a part of its own holding no picture: drawn by the layout, but
            # read for pictures alone it was never a node, and had no crop and no caption.
            chart = Chart(parse_chart(bytes(chart_part.blob)), chart_part)  # type: ignore[no-untyped-call]
            chart_id = stable_id(
                content_hash,
                canonical_locator("chart", {"body": body_index, "chart": chart_index}),
            )
            values = chart_metadata(chart)
            title = values.get("title")
            locator = SourceLocator(
                file_name=source.name,
                label=f"chart {body_index}.{chart_index}",
                relationship_id=relationship_id,
                source_part=str(chart_part.partname),
                metadata={"body_index": body_index, "chart_index": chart_index},
            )
            plotted: dict[str, JsonValue] = {
                "body_index": body_index, "chart_index": chart_index, **values,
            }
            nodes.append(
                GraphNode(
                    id=chart_id,
                    document_id=document_id,
                    type=NodeType.CHART,
                    label=title if isinstance(title, str) and title else f"Chart {chart_index}",
                    locator=locator,
                    metadata=plotted,
                )
            )
            evidence.append(
                EvidenceRecord(
                    id=stable_id(content_hash, f"evidence:chart:{body_index}:{chart_index}"),
                    document_id=document_id,
                    node_id=chart_id,
                    kind=EvidenceKind.SOURCE_METADATA,
                    locator=locator,
                    data=plotted,
                )
            )
            ordered_ids.append(chart_id)

    @staticmethod
    def _charts(paragraph: Paragraph, document: DocxDocument) -> Iterable[tuple[str, Any]]:
        for run in paragraph.runs:
            for found in run._element.iter(CHART_TAG):
                relationship_id = found.get(qn("r:id"))
                part = document.part.related_parts.get(relationship_id or "")
                if part is not None and hasattr(part, "blob"):
                    yield relationship_id, part

    @staticmethod
    def _extract_table(
        *,
        table: Table,
        body_index: int,
        source: Path,
        content_hash: str,
        document_id: str,
    ) -> tuple[GraphNode, EvidenceRecord]:
        rows = _grid(table)
        node_id = stable_id(
            content_hash, canonical_locator("table", {"body": body_index})
        )
        locator = SourceLocator(
            file_name=source.name,
            label=f"body table {body_index}",
            metadata={"body_index": body_index},
        )
        text = _grid_text(rows)
        node = GraphNode(
            id=node_id,
            document_id=document_id,
            type=NodeType.TABLE,
            label=f"Table {body_index}",
            locator=locator,
            text=text or None,
            content_hash=sha256_bytes(text.encode()),
            metadata={
                "body_index": body_index,
                "row_count": len(rows),
                "column_count": max((len(row) for row in rows), default=0),
            },
        )
        record = EvidenceRecord(
            id=stable_id(content_hash, f"evidence:table:{body_index}"),
            document_id=document_id,
            node_id=node_id,
            kind=EvidenceKind.SOURCE_VALUES,
            locator=locator,
            text=text or None,
            data={"rows": rows},
        )
        return node, record

    @staticmethod
    def _inline_images(
        paragraph: Paragraph, document: DocxDocument
    ) -> Iterable[tuple[str, Any]]:
        for run in paragraph.runs:
            for blip in run._element.iter(qn("a:blip")):
                relationship_id = blip.get(qn("r:embed"))
                if relationship_id:
                    image_part = document.part.related_parts.get(relationship_id)
                    if image_part is not None and hasattr(image_part, "blob"):
                        yield relationship_id, image_part

    @staticmethod
    def _text_boxes(paragraph: Paragraph) -> list[str]:
        """The words drawn in each text box this paragraph anchors, in document order.

        `document.paragraphs` walks the body and never enters a text box, so a report
        whose title block floats had no title in the graph at all.
        """
        found: list[str] = []
        for content in paragraph._p.iter(qn("w:txbxContent")):
            # Word writes a shape twice: the modern one and a VML fallback for old
            # readers, saying the same thing. Reading both duplicated every floating box.
            if any(node.tag == _MC_FALLBACK for node in content.iterancestors()):
                continue
            lines = [
                "".join(node.text or "" for node in block.iter(qn("w:t"))).strip()
                for block in content.findall(qn("w:p"))
            ]
            text = "\n".join(line for line in lines if line)
            if text:
                found.append(text)
        return found

    @staticmethod
    def _display_text(paragraph: Paragraph, raw: str) -> str:
        """Return visible text with a deterministic native-list prefix."""
        text = raw.strip()
        paragraph_xml = cast(Any, paragraph._p)
        num_id_nodes = paragraph_xml.xpath("./w:pPr/w:numPr/w:numId")
        if not num_id_nodes:
            return text
        num_id = num_id_nodes[0].get(qn("w:val"))
        level_nodes = paragraph_xml.xpath("./w:pPr/w:numPr/w:ilvl")
        ilvl = int(level_nodes[0].get(qn("w:val")) or 0) if level_nodes else 0
        if not num_id:
            return text
        # Preserve list presence in the searchable display text. Full numbering
        # format reconstruction will replace this fallback in a later extractor version.
        return f"{num_id}.{ilvl + 1} {text}" if text else f"{num_id}.{ilvl + 1}"

    @staticmethod
    def _extension(content_type: str) -> str:
        return {
            "image/jpeg": ".jpg",
            "image/png": ".png",
            "image/gif": ".gif",
            "image/tiff": ".tiff",
            "image/bmp": ".bmp",
            "image/x-emf": ".emf",
            "image/emf": ".emf",
            "image/x-wmf": ".wmf",
            "image/wmf": ".wmf",
        }.get(content_type, "")


def _grid(table: Table) -> list[list[str]]:
    """A table's text, one slot per grid column: a merged cell's words where it begins and
    nowhere else -- across (`gridSpan`) and down (`vMerge`) -- and the columns a row leaves
    out at its start or end (`gridBefore`, `gridAfter`) empty. Read off the XML, which holds
    whatever a converter wrote: a merge continued from no row above is a cell like any other."""
    rows: list[list[str]] = []
    for tr in _unwrapped(table._tbl, qn("w:tr")):
        row = [""] * tr.grid_before
        for tc in _unwrapped(tr, qn("w:tc")):
            continued = tc.vMerge == "continue" and bool(rows)
            row += ["" if continued else _cell_text(_Cell(tc, table))]
            row += [""] * (tc.grid_span - 1)
        rows.append(row + [""] * tr.grid_after)
    return rows


def _grid_text(rows: list[list[str]]) -> str:
    return "\n".join(" | ".join(row) for row in rows)


def _cell_text(cell: _Cell) -> str:
    """A cell's words in the order the cell holds them, a table nested in it row by row."""
    parts = [
        _grid_text(_grid(each)) if isinstance(each, Table)
        else _paragraph_text(each)
        for each in _blocks(cell._tc, cell)
    ]
    return "\n".join(part for part in parts if part.strip()).strip()


def _blocks(element: Any, parent: Any) -> Iterator[Paragraph | Table]:
    """The paragraphs and tables `element` holds, in order, those inside a content control
    (`w:sdt`) included: a template keeps its cover page in one, and the layout counts each
    paragraph inside it as a body element of its own."""
    for child in _unwrapped(element, qn("w:p"), qn("w:tbl")):
        yield Paragraph(child, parent) if child.tag == qn("w:p") else Table(child, parent)


def _unwrapped(element: Any, *tags: str) -> Iterator[Any]:
    """`element`'s children with one of `tags`, looking through content controls -- around
    paragraphs, a table's rows (a repeating section) or a row's cells alike."""
    for child in element.iterchildren():
        if child.tag in tags:
            yield child
        elif child.tag == qn("w:sdt"):
            content = child.find(qn("w:sdtContent"))
            if content is not None:
                yield from _unwrapped(content, *tags)


_MC_FALLBACK = "{http://schemas.openxmlformats.org/markup-compatibility/2006}Fallback"

_NOT_THE_PARAGRAPHS_WORDS = {
    qn("w:del"),
    qn("w:moveFrom"),
    qn("w:txbxContent"),
    qn("w:rt"),
    _MC_FALLBACK,
}
"""What a tracked change deleted or moved away, a text box (a unit of its own), a ruby's
reading guide, and the copy Word writes for old readers."""


def _paragraph_text(paragraph: Paragraph) -> str:
    """A paragraph's words, those a tracked insertion, a simple field or an inline content
    control hold one level down included -- `Paragraph.text` reads only its own runs."""
    own = paragraph._p
    return "".join(
        run.text
        for run in own.iter(qn("w:r"))
        if not any(
            ancestor.tag in _NOT_THE_PARAGRAPHS_WORDS
            for ancestor in _ancestors_below(run, own)
        )
    )


def _ancestors_below(element: Any, top: Any) -> Iterator[Any]:
    for ancestor in element.iterancestors():
        if ancestor is top:
            return
        yield ancestor
