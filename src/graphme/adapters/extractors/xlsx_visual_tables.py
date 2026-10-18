"""Native detection of bordered, non-formal visual tables in XLSX sheets.

The detector intentionally uses workbook-native structure only.  It does not OCR or
interpret business labels: borders establish candidates, while merges, geometry,
cell value types, and drawing footprints describe their internal layout.
"""

from __future__ import annotations

from collections import Counter, defaultdict, deque
from collections.abc import Iterable, Sequence
from dataclasses import dataclass
from typing import Literal

from openpyxl.cell.cell import Cell, MergedCell
from openpyxl.drawing.spreadsheet_drawing import OneCellAnchor, TwoCellAnchor
from openpyxl.styles.borders import Side
from openpyxl.utils.cell import get_column_letter, range_boundaries
from openpyxl.worksheet.formula import ArrayFormula, DataTableFormula
from openpyxl.worksheet.worksheet import Worksheet

ValueKind = Literal["empty", "text", "number", "boolean", "date", "formula", "mixed"]
GroupRole = Literal["stub", "repeated", "field"]

PERIMETER_GAP = 0.1
"""How much of one side of a table's border may be missing and still read as a border."""

AXIS_MARGIN = 64
"""How far past the last anchored cell the row and column sizes are still measured. A
one-cell anchor states a size, not an end cell, so a picture can reach past everything the
sheet anchors; beyond this the axis carries on in the last size it measured."""

MAX_COLUMNS = 16384
MAX_ROWS = 1048576
"""What a sheet can hold, which is where measuring its axes has to stop."""

_EMU_PER_PIXEL = 9525.0
_DEFAULT_COLUMN_WIDTH = 13.0
_DEFAULT_ROW_HEIGHT = 15.0


@dataclass(frozen=True)
class GridRectangle:
    min_col: int
    min_row: int
    max_col: int
    max_row: int

    @property
    def ref(self) -> str:
        return (
            f"{get_column_letter(self.min_col)}{self.min_row}:"
            f"{get_column_letter(self.max_col)}{self.max_row}"
        )

    def contains(self, col: int, row: int) -> bool:
        return self.min_col <= col <= self.max_col and self.min_row <= row <= self.max_row

    def overlaps(self, other: GridRectangle) -> bool:
        return (
            self.min_col <= other.max_col
            and self.max_col >= other.min_col
            and self.min_row <= other.max_row
            and self.max_row >= other.min_row
        )


@dataclass(frozen=True)
class DrawingFootprint:
    """A drawing's physical footprint and conservative touched-cell rectangle."""

    index: int
    rectangle: GridRectangle
    left_px: float
    top_px: float
    right_px: float
    bottom_px: float


@dataclass(frozen=True)
class ColumnGroup:
    rectangle: GridRectangle
    role: GroupRole
    value_kind: ValueKind
    repeated_index: int | None = None


@dataclass(frozen=True)
class RowBand:
    rectangle: GridRectangle
    value_kind: ValueKind


@dataclass(frozen=True)
class ImageSlot:
    rectangle: GridRectangle
    drawing_indices: tuple[int, ...]


@dataclass(frozen=True)
class VisualTableCandidate:
    sheet: str
    rectangle: GridRectangle
    values: tuple[tuple[object | None, ...], ...]
    column_groups: tuple[ColumnGroup, ...]
    row_bands: tuple[RowBand, ...]
    image_slots: tuple[ImageSlot, ...]
    drawings: tuple[DrawingFootprint, ...]
    merges: tuple[GridRectangle, ...]
    signals: dict[str, int | bool]

    @property
    def ref(self) -> str:
        return self.rectangle.ref


def drawing_footprints(sheet: Worksheet) -> tuple[DrawingFootprint, ...]:
    """Return complete native drawing footprints for a worksheet.

    Two-cell anchors use both markers and their offsets. One-cell anchors use the
    start marker plus the drawing's intrinsic dimensions, which is the best native
    approximation available without rendering Excel.

    The axes are measured as far as the anchors reach and grow from there on demand. A
    sheet that declares a whole column as its dimension -- Excel writes that of its own
    accord -- otherwise had a million row heights measured, one openpyxl row object at a
    time, to place a handful of pictures.
    """

    anchors = [
        drawing.anchor
        for drawing in getattr(sheet, "_images", ())
        if isinstance(drawing.anchor, (OneCellAnchor, TwoCellAnchor))
    ]
    if not anchors:
        return ()
    columns = _AxisGeometry(
        _column_pixels(sheet, col)
        for col in range(1, _measured(anchors, "col", MAX_COLUMNS) + 1)
    )
    rows = _AxisGeometry(
        _row_pixels(sheet, row) for row in range(1, _measured(anchors, "row", MAX_ROWS) + 1)
    )
    result: list[DrawingFootprint] = []
    for index, drawing in enumerate(getattr(sheet, "_images", ())):
        anchor = drawing.anchor
        if not isinstance(anchor, (OneCellAnchor, TwoCellAnchor)):
            continue
        left = columns.position(anchor._from.col) + anchor._from.colOff / _EMU_PER_PIXEL
        top = rows.position(anchor._from.row) + anchor._from.rowOff / _EMU_PER_PIXEL
        if isinstance(anchor, TwoCellAnchor):
            right = columns.position(anchor.to.col) + anchor.to.colOff / _EMU_PER_PIXEL
            bottom = rows.position(anchor.to.row) + anchor.to.rowOff / _EMU_PER_PIXEL
        else:
            right = left + float(getattr(drawing, "width", 0.0))
            bottom = top + float(getattr(drawing, "height", 0.0))
        if right <= left or bottom <= top:
            continue
        rectangle = GridRectangle(
            min_col=columns.index_at(left) + 1,
            min_row=rows.index_at(top) + 1,
            max_col=columns.index_at(max(left, right - 1e-6)) + 1,
            max_row=rows.index_at(max(top, bottom - 1e-6)) + 1,
        )
        result.append(
            DrawingFootprint(index, rectangle, left, top, right, bottom)
        )
    return tuple(result)


def _measured(
    anchors: Sequence[OneCellAnchor | TwoCellAnchor], axis: str, limit: int
) -> int:
    """How far along one axis the sizes are read: the last anchor, plus the margin.

    Both markers of every anchor count, and the sheet's own end is the ceiling -- Excel
    has no column past its last one to measure.
    """
    marked: list[int] = [int(getattr(anchor._from, axis)) for anchor in anchors]
    marked += [
        int(getattr(anchor.to, axis)) for anchor in anchors if isinstance(anchor, TwoCellAnchor)
    ]
    return min(limit, max(marked) + 1 + AXIS_MARGIN)


def detect_visual_tables(
    sheet: Worksheet, cells: Sequence[Cell] | None = None
) -> tuple[VisualTableCandidate, ...]:
    """Detect complete bordered visual-table rectangles conservatively.

    `cells` is `written_cells(sheet)` when the caller already has it."""

    horizontal, vertical = _border_segments(written_cells(sheet) if cells is None else cells)
    formal = tuple(GridRectangle(*range_boundaries(table.ref)) for table in sheet.tables.values())
    drawings = drawing_footprints(sheet)
    sheet_merges = tuple(
        GridRectangle(merge.min_col, merge.min_row, merge.max_col, merge.max_row)
        for merge in sheet.merged_cells.ranges
    )
    candidates: list[VisualTableCandidate] = []
    for rectangle in _component_rectangles(horizontal, vertical):
        if rectangle.max_col <= rectangle.min_col or rectangle.max_row <= rectangle.min_row:
            continue
        if not _complete_perimeter(rectangle, horizontal, vertical):
            continue
        if any(rectangle.overlaps(table) for table in formal):
            continue
        internal_h = _full_horizontal_rules(rectangle, horizontal)
        internal_v = _full_vertical_rules(rectangle, vertical)
        merges = _contained_merges(sheet_merges, rectangle)
        populated, styled, kinds = _scan(sheet, rectangle)
        repeated = _repeated_column_spans(merges, rectangle)
        # Outer boxes and decorated prose are deliberately rejected. A visual table
        # needs internal rules plus either repeated merge topology or enough values.
        # Fill/font/alignment only help a nearly sufficient value signal; they can
        # never create a candidate without the structural border requirements.
        if not internal_h and not internal_v:
            continue
        if not repeated and populated < 4 and not (populated >= 3 and styled >= 3):
            continue
        table_drawings = tuple(d for d in drawings if d.rectangle.overlaps(rectangle))
        values = _values(sheet, sheet_merges, rectangle)
        candidates.append(
            VisualTableCandidate(
                sheet=sheet.title,
                rectangle=rectangle,
                values=values,
                column_groups=_column_groups(kinds, rectangle, repeated),
                row_bands=_row_bands(kinds, rectangle, internal_h),
                image_slots=_image_slots(merges, table_drawings),
                drawings=table_drawings,
                merges=merges,
                signals={
                    "populated_cells": populated,
                    "styled_cells": styled,
                    "internal_horizontal_rules": len(internal_h),
                    "internal_vertical_rules": len(internal_v),
                    "repeated_column_topology": bool(repeated),
                },
            )
        )
    return tuple(
        sorted(
            candidates,
            key=lambda item: (item.rectangle.min_row, item.rectangle.min_col),
        )
    )


def _border_segments(cells: Iterable[Cell]) -> tuple[set[tuple[int, int]], set[tuple[int, int]]]:
    """Return horizontal (grid-row, col) and vertical (row, grid-col) unit segments.

    Only the cells the file actually wrote are read. Walking the declared dimension
    instead built a cell object for every position inside it, and a sheet that declares a
    whole column -- Excel writes that of its own accord -- made a million of them per
    scan; a cell with a border is a cell the file wrote, so nothing is missed.
    """

    horizontal: set[tuple[int, int]] = set()
    vertical: set[tuple[int, int]] = set()
    for cell in cells:
        r, c = cell.row, cell.column
        if _ruled(cell.border.top):
            horizontal.add((r, c))
        if _ruled(cell.border.bottom):
            horizontal.add((r + 1, c))
        if _ruled(cell.border.left):
            vertical.add((r, c))
        if _ruled(cell.border.right):
            vertical.add((r, c + 1))
    return horizontal, vertical


def is_formula(value: object) -> bool:
    """Whether a cell's value is a formula: openpyxl reads an array or data-table formula
    as an object of its own, not as text starting with `=`."""
    return isinstance(value, (ArrayFormula, DataTableFormula)) or (
        isinstance(value, str) and value.startswith("=")
    )


def written_cells(sheet: Worksheet) -> list[Cell]:
    """The cells the sheet holds, in reading order, whatever its dimension declares.

    Walking the declared dimension builds a cell object for every position inside it, and
    Excel writes dimensions of its own accord that reach a whole column past the last
    thing on the sheet -- a million cells made and thrown away, per scan. A cell that
    holds anything, a value or a border, is a cell the file wrote.
    """
    written: dict[tuple[int, int], Cell] | None = getattr(sheet, "_cells", None)
    if written is None:  # a worksheet that keeps no cell map answers the long way
        return [cell for row in sheet.iter_rows() for cell in row]
    return [cell for _, cell in sorted(written.items())]


def detect_data_blocks(
    cells: Iterable[Cell], taken: Sequence[GridRectangle], *, min_rows: int
) -> tuple[str, ...]:
    """Ranges where `min_rows` rows or more run unbroken, each holding two values or more.

    `cells` are the sheet's `written_cells`. `taken` is what is already a table; its cells
    neither start a block nor join one, and a block that would cross one is not made. A
    block's columns are those filled in at least half its rows, so a note written beside
    one row does not widen it.
    """
    by_row: dict[int, list[int]] = defaultdict(list)
    for cell in cells:
        if isinstance(cell, MergedCell) or cell.value is None or not str(cell.value).strip():
            continue
        if any(box.contains(cell.column, cell.row) for box in taken):
            continue
        by_row[cell.row].append(cell.column)

    runs: list[list[int]] = []
    for row in sorted(by_row):
        if len(by_row[row]) < 2:
            continue
        if runs and runs[-1][-1] == row - 1:
            runs[-1].append(row)
        else:
            runs.append([row])

    blocks: list[str] = []
    for run in runs:
        if len(run) < min_rows:
            continue
        counts = Counter(col for row in run for col in by_row[row])
        cols = [col for col, seen in counts.items() if seen * 2 >= len(run)]
        if len(cols) < 2:
            continue
        block = GridRectangle(min(cols), run[0], max(cols), run[-1])
        if any(block.overlaps(other) for other in taken):
            continue
        blocks.append(block.ref)
    return tuple(blocks)


def _ruled(side: Side | None) -> bool:
    """Whether a cell edge is drawn. A side the file never wrote comes back as nothing."""
    return side is not None and side.style is not None


def _component_rectangles(
    horizontal: set[tuple[int, int]], vertical: set[tuple[int, int]]
) -> tuple[GridRectangle, ...]:
    edges: dict[tuple[int, int], set[tuple[int, int]]] = defaultdict(set)
    for grid_row, col in horizontal:
        left, right = (grid_row, col), (grid_row, col + 1)
        edges[left].add(right)
        edges[right].add(left)
    for row, grid_col in vertical:
        top, bottom = (row, grid_col), (row + 1, grid_col)
        edges[top].add(bottom)
        edges[bottom].add(top)
    unseen = set(edges)
    rectangles: list[GridRectangle] = []
    while unseen:
        start = unseen.pop()
        queue = deque([start])
        vertices = [start]
        while queue:
            vertex = queue.popleft()
            for adjacent in edges[vertex]:
                if adjacent in unseen:
                    unseen.remove(adjacent)
                    queue.append(adjacent)
                    vertices.append(adjacent)
        grid_rows = [vertex[0] for vertex in vertices]
        grid_cols = [vertex[1] for vertex in vertices]
        rectangles.append(
            GridRectangle(min(grid_cols), min(grid_rows), max(grid_cols) - 1, max(grid_rows) - 1)
        )
    return tuple(rectangles)


def _complete_perimeter(
    rectangle: GridRectangle,
    horizontal: set[tuple[int, int]],
    vertical: set[tuple[int, int]],
) -> bool:
    columns = range(rectangle.min_col, rectangle.max_col + 1)
    rows = range(rectangle.min_row, rectangle.max_row + 1)
    sides = (
        [(rectangle.min_row, col) in horizontal for col in columns],
        [(rectangle.max_row + 1, col) in horizontal for col in columns],
        [(row, rectangle.min_col) in vertical for row in rows],
        [(row, rectangle.max_col + 1) in vertical for row in rows],
    )
    return all(_drawn_side(side) for side in sides)


def _drawn_side(side: Sequence[bool]) -> bool:
    """Whether one side of the rectangle reads as drawn, unit by unit.

    A side had to be unbroken, and one cell of a long bottom row that its author never
    formatted was enough to reject the whole table -- it was then read out as loose text,
    a cell at a time, with its rows and columns gone. A gap is allowed in proportion to
    the side, so a short box still needs every unit of it.
    """
    return len(side) > 0 and side.count(False) <= int(len(side) * PERIMETER_GAP)


def _full_horizontal_rules(
    rectangle: GridRectangle, horizontal: set[tuple[int, int]]
) -> tuple[int, ...]:
    return tuple(
        grid_row
        for grid_row in range(rectangle.min_row + 1, rectangle.max_row + 1)
        if all(
            (grid_row, col) in horizontal
            for col in range(rectangle.min_col, rectangle.max_col + 1)
        )
    )


def _full_vertical_rules(
    rectangle: GridRectangle, vertical: set[tuple[int, int]]
) -> tuple[int, ...]:
    return tuple(
        grid_col
        for grid_col in range(rectangle.min_col + 1, rectangle.max_col + 1)
        if all(
            (row, grid_col) in vertical
            for row in range(rectangle.min_row, rectangle.max_row + 1)
        )
    )


def _contained_merges(
    merges: Sequence[GridRectangle], rectangle: GridRectangle
) -> tuple[GridRectangle, ...]:
    return tuple(
        merge
        for merge in merges
        if (
            rectangle.min_col <= merge.min_col <= merge.max_col <= rectangle.max_col
            and rectangle.min_row <= merge.min_row <= merge.max_row <= rectangle.max_row
        )
    )


def _repeated_column_spans(
    merges: Sequence[GridRectangle], rectangle: GridRectangle
) -> tuple[tuple[int, int], ...]:
    patterns: Counter[tuple[int, tuple[int, ...]]] = Counter()
    spans_for_pattern: dict[tuple[int, tuple[int, ...]], tuple[tuple[int, int], ...]] = {}
    by_rows: dict[tuple[int, int], list[tuple[int, int]]] = defaultdict(list)
    for merge in merges:
        if merge.max_col > merge.min_col:
            by_rows[(merge.min_row, merge.max_row)].append((merge.min_col, merge.max_col))
    for spans in by_rows.values():
        ordered = tuple(sorted(spans))
        widths = Counter(end - start + 1 for start, end in ordered)
        for width, count in widths.items():
            repeated = tuple(span for span in ordered if span[1] - span[0] + 1 == width)
            if count >= 2:
                gaps = tuple(
                    repeated[i + 1][0] - repeated[i][1] - 1
                    for i in range(len(repeated) - 1)
                )
                key = (width, gaps)
                patterns[key] += len(repeated)
                spans_for_pattern[key] = repeated
    if not patterns:
        return ()
    best = max(patterns, key=lambda key: (patterns[key], key[0] * len(spans_for_pattern[key])))
    result = spans_for_pattern[best]
    # Prefer a topology that covers a meaningful part of the candidate width.
    covered_width = sum(end - start + 1 for start, end in result)
    if covered_width < (rectangle.max_col - rectangle.min_col + 1) // 2:
        return ()
    return result


def _column_groups(
    kinds: dict[tuple[int, int], ValueKind],
    rectangle: GridRectangle,
    repeated: Sequence[tuple[int, int]],
) -> tuple[ColumnGroup, ...]:
    groups: list[ColumnGroup] = []
    cursor = rectangle.min_col
    for index, (start, end) in enumerate(repeated):
        if cursor < start:
            role: GroupRole = "stub" if cursor == rectangle.min_col else "field"
            area = GridRectangle(cursor, rectangle.min_row, start - 1, rectangle.max_row)
            groups.append(ColumnGroup(area, role, _value_kind(kinds, area)))
        area = GridRectangle(start, rectangle.min_row, end, rectangle.max_row)
        groups.append(ColumnGroup(area, "repeated", _value_kind(kinds, area), index))
        cursor = end + 1
    if cursor <= rectangle.max_col:
        role = "stub" if not groups else "field"
        area = GridRectangle(cursor, rectangle.min_row, rectangle.max_col, rectangle.max_row)
        groups.append(ColumnGroup(area, role, _value_kind(kinds, area)))
    if not groups:
        for col in range(rectangle.min_col, rectangle.max_col + 1):
            area = GridRectangle(col, rectangle.min_row, col, rectangle.max_row)
            groups.append(ColumnGroup(area, "field", _value_kind(kinds, area)))
    return tuple(groups)


def _row_bands(
    kinds: dict[tuple[int, int], ValueKind], rectangle: GridRectangle, internal_rules: Sequence[int]
) -> tuple[RowBand, ...]:
    starts = (rectangle.min_row, *internal_rules)
    ends = (*[rule - 1 for rule in internal_rules], rectangle.max_row)
    return tuple(
        RowBand(
            area := GridRectangle(rectangle.min_col, start, rectangle.max_col, end),
            _value_kind(kinds, area),
        )
        for start, end in zip(starts, ends, strict=True)
    )


def _image_slots(
    merges: Sequence[GridRectangle], drawings: Sequence[DrawingFootprint]
) -> tuple[ImageSlot, ...]:
    # Large merged regions are explicit layout slots, including deliberately empty
    # ones. A one-row merged heading is not an image slot.
    slots = [
        merge
        for merge in merges
        if merge.max_col > merge.min_col and merge.max_row > merge.min_row
    ]
    return tuple(
        ImageSlot(
            slot,
            tuple(drawing.index for drawing in drawings if drawing.rectangle.overlaps(slot)),
        )
        for slot in sorted(slots, key=lambda item: (item.min_row, item.min_col))
    )


def _scan(
    sheet: Worksheet, rectangle: GridRectangle
) -> tuple[int, int, dict[tuple[int, int], ValueKind]]:
    """How many cells in the rectangle hold a value, how many are styled, and the kind of
    each value, read in one pass over the cells the file wrote: `sheet.cell()` would insert
    an empty one for each miss. A position the file left out has the sheet's default style,
    which is what an empty cell made for it would have."""
    blank_styled = _styled(Cell(sheet))
    populated = styled = 0
    kinds: dict[tuple[int, int], ValueKind] = {}
    for row in range(rectangle.min_row, rectangle.max_row + 1):
        for col in range(rectangle.min_col, rectangle.max_col + 1):
            cell = sheet._cells.get((row, col))
            if cell is None:
                styled += blank_styled
                continue
            styled += _styled(cell)
            if cell.value is None:
                continue
            populated += 1
            if not isinstance(cell, MergedCell):
                kinds[(row, col)] = _kind_of(cell)
    return populated, styled, kinds


def _styled(cell: Cell | MergedCell) -> bool:
    has_fill = cell.fill.fill_type is not None
    has_font = cell.font.bold or cell.font.italic or cell.font.sz is not None
    has_alignment = any(
        value is not None
        for value in (
            cell.alignment.horizontal,
            cell.alignment.vertical,
            cell.alignment.wrap_text,
        )
    )
    return bool(has_fill or has_font or has_alignment)


def _values(
    sheet: Worksheet, merges: Sequence[GridRectangle], rectangle: GridRectangle
) -> tuple[tuple[object | None, ...], ...]:
    def value_at(row: int, col: int) -> object | None:
        cell = sheet._cells.get((row, col))
        return None if cell is None else cell.value

    anchors: dict[tuple[int, int], object | None] = {}
    for merge in merges:
        if not merge.overlaps(rectangle):
            continue
        value = value_at(merge.min_row, merge.min_col)
        for row in range(merge.min_row, merge.max_row + 1):
            for col in range(merge.min_col, merge.max_col + 1):
                anchors[(row, col)] = value
    return tuple(
        tuple(
            anchors.get((row, col), value_at(row, col))
            for col in range(rectangle.min_col, rectangle.max_col + 1)
        )
        for row in range(rectangle.min_row, rectangle.max_row + 1)
    )


def _kind_of(cell: Cell) -> ValueKind:
    if cell.data_type == "f" or is_formula(cell.value):
        return "formula"
    if cell.is_date:
        return "date"
    if isinstance(cell.value, bool):
        return "boolean"
    if isinstance(cell.value, (int, float)):
        return "number"
    return "text"


def _value_kind(
    kinds: dict[tuple[int, int], ValueKind], rectangle: GridRectangle
) -> ValueKind:
    found = {kind for (row, col), kind in kinds.items() if rectangle.contains(col, row)}
    if not found:
        return "empty"
    if len(found) == 1:
        return next(iter(found))
    return "mixed"


class _AxisGeometry:
    def __init__(self, sizes: Iterable[float]) -> None:
        self._positions = [0.0]
        for size in sizes:
            self._positions.append(self._positions[-1] + size)

    def position(self, zero_based_index: int) -> float:
        while zero_based_index >= len(self._positions):
            last_size = self._positions[-1] - self._positions[-2]
            self._positions.append(self._positions[-1] + last_size)
        return self._positions[zero_based_index]

    def index_at(self, position: float) -> int:
        while position >= self._positions[-1]:
            last_size = self._positions[-1] - self._positions[-2]
            self._positions.append(self._positions[-1] + last_size)
        low, high = 0, len(self._positions) - 1
        while low + 1 < high:
            middle = (low + high) // 2
            if self._positions[middle] <= position:
                low = middle
            else:
                high = middle
        return low


def _column_pixels(sheet: Worksheet, col: int) -> float:
    width = sheet.column_dimensions[get_column_letter(col)].width
    if width is None:
        width = sheet.sheet_format.defaultColWidth or _DEFAULT_COLUMN_WIDTH
    # Excel's documented character-width approximation at 96 DPI.
    return float(int(width * 7 + 5))


def _row_pixels(sheet: Worksheet, row: int) -> float:
    height = sheet.row_dimensions[row].height
    if height is None:
        height = sheet.sheet_format.defaultRowHeight or _DEFAULT_ROW_HEIGHT
    return float(height) * 96.0 / 72.0
