from __future__ import annotations

import hashlib
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from openpyxl.utils.cell import get_column_letter

from graphme.application.file_titles import OFFICE_REL, relations, rels_part

_XDR = "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing"
_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
_R = OFFICE_REL
_MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"


def _local(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def _int(value: str | None) -> int | None:
    return int(value) if value is not None else None


def _hash(payload: bytes) -> str:
    return hashlib.sha256(payload).hexdigest()


def _relationships(archive: ZipFile, part: str) -> dict[str, tuple[str, str]]:
    rels_name = rels_part(part)
    if rels_name not in archive.namelist():
        return {}
    return {rid: (target, kind) for _, rid, kind, target in relations(archive, rels_name)}


@dataclass(frozen=True)
class AnchorMarker:
    col: int | None = None
    col_off: int | None = None
    row: int | None = None
    row_off: int | None = None
    x: int | None = None
    y: int | None = None
    cx: int | None = None
    cy: int | None = None

    def as_dict(self) -> dict[str, int | None]:
        return {
            "col": self.col,
            "col_off": self.col_off,
            "row": self.row,
            "row_off": self.row_off,
            "x": self.x,
            "y": self.y,
            "cx": self.cx,
            "cy": self.cy,
        }


@dataclass(frozen=True)
class DrawingObject:
    object_id: str
    worksheet: str
    worksheet_part: str
    drawing_part: str
    drawing_relationship_id: str
    xml_index: int
    z_index: int
    kind: str
    anchor_kind: str
    anchor_from: AnchorMarker | None
    anchor_to: AnchorMarker | None
    extent: AnchorMarker | None
    edit_as: str | None
    nv_id: int | None
    name: str | None
    group_path: tuple[int, ...]
    transform: dict[str, Any]
    properties: dict[str, Any]
    text_runs: tuple[dict[str, Any], ...]
    media_relationship_id: str | None
    media_target: str | None
    media_hash: str | None
    media_bytes: bytes | None = field(repr=False)
    raw_xml_hash: str = ""

    @property
    def is_picture(self) -> bool:
        return self.kind == "pic" and self.media_bytes is not None


@dataclass(frozen=True)
class DrawingManifest:
    workbook_hash: str
    objects: tuple[DrawingObject, ...]
    warnings: tuple[str, ...] = ()


@dataclass(frozen=True)
class DrawingComponent:
    component_id: str
    worksheet: str
    drawing_part: str
    members: tuple[DrawingObject, ...]
    source_fidelity: str
    warning: str | None = None

    @property
    def pictures(self) -> tuple[DrawingObject, ...]:
        return tuple(member for member in self.members if member.is_picture)

def _marker(element: ET.Element | None) -> AnchorMarker | None:
    if element is None:
        return None
    values = {_local(child.tag): _int(child.text) for child in element}
    return AnchorMarker(
        col=values.get("col"),
        col_off=values.get("colOff"),
        row=values.get("row"),
        row_off=values.get("rowOff"),
        x=_int(element.get("x")),
        y=_int(element.get("y")),
        cx=_int(element.get("cx")),
        cy=_int(element.get("cy")),
    )


def _xfrm(element: ET.Element) -> dict[str, Any]:
    xfrm = element.find(f".//{{{_A}}}xfrm")
    if xfrm is None:
        xfrm = element.find(f".//{{{_XDR}}}xfrm")
    if xfrm is None:
        return {}
    result: dict[str, Any] = {
        "rotation": _int(xfrm.get("rot")),
        "flip_h": xfrm.get("flipH") == "1",
        "flip_v": xfrm.get("flipV") == "1",
    }
    for name in ("off", "ext", "chOff", "chExt"):
        child = xfrm.find(f"{{{_A}}}{name}")
        if child is not None:
            result[name] = {key: _int(value) for key, value in child.attrib.items()}
    return result


def _color(element: ET.Element | None) -> dict[str, Any] | None:
    if element is None:
        return None
    color = next(iter(element), None)
    if color is None:
        return {"kind": _local(element.tag)}
    result: dict[str, Any] = {"kind": _local(color.tag), "value": color.get("val")}
    alpha = color.find(f"{{{_A}}}alpha")
    if alpha is not None:
        result["alpha"] = _int(alpha.get("val"))
    return result


def _shape_properties(element: ET.Element) -> dict[str, Any]:
    connector_properties = next(
        (
            candidate
            for candidate in element.iter()
            if _local(candidate.tag) == "cNvCxnSpPr"
        ),
        None,
    )
    connections: dict[str, dict[str, int | None]] = {}
    if connector_properties is not None:
        for endpoint in connector_properties:
            name = _local(endpoint.tag)
            if name in {"stCxn", "endCxn"}:
                connections[name] = {
                    "id": _int(endpoint.get("id")),
                    "index": _int(endpoint.get("idx")),
                }
    sp_pr = element.find(f"{{{_XDR}}}spPr")
    if sp_pr is None:
        sp_pr = element.find(f"{{{_A}}}spPr")
    if sp_pr is None:
        return {"connections": connections} if connections else {}
    geometry = sp_pr.find(f"{{{_A}}}prstGeom")
    custom = sp_pr.find(f"{{{_A}}}custGeom")
    fill = next(
        (
            child
            for child in sp_pr
            if _local(child.tag) in {"solidFill", "gradFill", "pattFill", "noFill"}
        ),
        None,
    )
    line = sp_pr.find(f"{{{_A}}}ln")
    result: dict[str, Any] = {
        "geometry": geometry.get("prst")
        if geometry is not None
        else "custom"
        if custom is not None
        else None,
        "fill": _color(fill),
    }
    if line is not None:
        line_fill = next(
            (
                child
                for child in line
                if _local(child.tag) in {"solidFill", "gradFill", "pattFill", "noFill"}
            ),
            None,
        )
        head = line.find(f"{{{_A}}}headEnd")
        tail = line.find(f"{{{_A}}}tailEnd")
        result["line"] = {
            "width": _int(line.get("w")),
            "fill": _color(line_fill),
            "head": dict(head.attrib) if head is not None else None,
            "tail": dict(tail.attrib) if tail is not None else None,
        }
    src_rect = element.find(f".//{{{_A}}}srcRect")
    if src_rect is not None:
        result["crop"] = {key: _int(value) for key, value in src_rect.attrib.items()}
    if connections:
        result["connections"] = connections
    return result


def _text_runs(element: ET.Element) -> tuple[dict[str, Any], ...]:
    runs: list[dict[str, Any]] = []
    for paragraph_index, paragraph in enumerate(element.findall(f".//{{{_A}}}p")):
        for run_index, run in enumerate(paragraph.findall(f"{{{_A}}}r")):
            text = run.findtext(f"{{{_A}}}t", default="")
            properties = run.find(f"{{{_A}}}rPr")
            runs.append(
                {
                    "paragraph": paragraph_index,
                    "run": run_index,
                    "text": text,
                    "properties": dict(properties.attrib) if properties is not None else {},
                }
            )
        field_text = paragraph.findtext(f"{{{_A}}}fld/{{{_A}}}t")
        if field_text:
            runs.append(
                {
                    "paragraph": paragraph_index,
                    "run": len(runs),
                    "text": field_text,
                    "properties": {},
                }
            )
    return tuple(runs)


def _nonvisual(element: ET.Element) -> tuple[int | None, str | None]:
    for candidate in element.iter():
        if _local(candidate.tag) == "cNvPr":
            return _int(candidate.get("id")), candidate.get("name")
    return None, None


def _kind(element: ET.Element) -> str:
    name = _local(element.tag)
    return name if name in {"pic", "sp", "cxnSp", "grpSp", "graphicFrame"} else "unknown"


def _walk_objects(
    *,
    element: ET.Element,
    archive: ZipFile,
    workbook_hash: str,
    worksheet: str,
    worksheet_part: str,
    drawing_part: str,
    drawing_rid: str,
    drawing_rels: dict[str, tuple[str, str]],
    xml_index: int,
    z_index: int,
    anchor_kind: str,
    anchor_from: AnchorMarker | None,
    anchor_to: AnchorMarker | None,
    extent: AnchorMarker | None,
    edit_as: str | None,
    group_path: tuple[int, ...],
) -> list[DrawingObject]:
    kind = _kind(element)
    nv_id, name = _nonvisual(element)
    media_rid: str | None = None
    media_target: str | None = None
    media_bytes: bytes | None = None
    blip = element.find(f".//{{{_A}}}blip")
    if blip is not None:
        media_rid = blip.get(f"{{{_R}}}embed") or blip.get(f"{{{_R}}}link")
        if media_rid in drawing_rels:
            target, _ = drawing_rels[media_rid]
            media_target = target
            if target in archive.namelist():
                media_bytes = archive.read(target)
    raw = ET.tostring(element, encoding="utf-8")
    identity = ":".join(
        (
            workbook_hash,
            worksheet,
            drawing_part,
            str(xml_index),
            str(nv_id),
            ".".join(map(str, group_path)),
        )
    )
    item = DrawingObject(
        object_id=hashlib.sha256(identity.encode()).hexdigest(),
        worksheet=worksheet,
        worksheet_part=worksheet_part,
        drawing_part=drawing_part,
        drawing_relationship_id=drawing_rid,
        xml_index=xml_index,
        z_index=z_index,
        kind=kind,
        anchor_kind=anchor_kind,
        anchor_from=anchor_from,
        anchor_to=anchor_to,
        extent=extent,
        edit_as=edit_as,
        nv_id=nv_id,
        name=name,
        group_path=group_path,
        transform=_xfrm(element),
        properties=_shape_properties(element),
        text_runs=_text_runs(element),
        media_relationship_id=media_rid,
        media_target=media_target,
        media_hash=_hash(media_bytes) if media_bytes is not None else None,
        media_bytes=media_bytes,
        raw_xml_hash=_hash(raw),
    )
    result = [item]
    if kind == "grpSp":
        child_number = 0
        for child in element:
            if _local(child.tag) in {"nvGrpSpPr", "grpSpPr"}:
                continue
            result.extend(
                _walk_objects(
                    element=child,
                    archive=archive,
                    workbook_hash=workbook_hash,
                    worksheet=worksheet,
                    worksheet_part=worksheet_part,
                    drawing_part=drawing_part,
                    drawing_rid=drawing_rid,
                    drawing_rels=drawing_rels,
                    xml_index=xml_index,
                    z_index=z_index,
                    anchor_kind=anchor_kind,
                    anchor_from=anchor_from,
                    anchor_to=anchor_to,
                    extent=extent,
                    edit_as=edit_as,
                    group_path=(*group_path, child_number),
                )
            )
            child_number += 1
    return result


def parse_xlsx_drawings(source: Path, *, workbook_hash: str | None = None) -> DrawingManifest:
    """Parse the OOXML drawing graph without relying on openpyxl private image lists."""
    payload_hash = workbook_hash or _hash(source.read_bytes())
    objects: list[DrawingObject] = []
    warnings: list[str] = []
    with ZipFile(source) as archive:
        workbook_part = "xl/workbook.xml"
        workbook = ET.fromstring(archive.read(workbook_part))
        workbook_rels = _relationships(archive, workbook_part)
        for sheet in workbook.findall(f".//{{{_MAIN}}}sheet"):
            title = sheet.get("name", "")
            sheet_rid = sheet.get(f"{{{_R}}}id")
            if not sheet_rid or sheet_rid not in workbook_rels:
                continue
            worksheet_part = workbook_rels[sheet_rid][0]
            if worksheet_part not in archive.namelist():
                continue
            worksheet = ET.fromstring(archive.read(worksheet_part))
            sheet_rels = _relationships(archive, worksheet_part)
            for drawing_ref in worksheet.findall(f"{{{_MAIN}}}drawing"):
                drawing_rid = drawing_ref.get(f"{{{_R}}}id")
                if not drawing_rid or drawing_rid not in sheet_rels:
                    continue
                drawing_part = sheet_rels[drawing_rid][0]
                if drawing_part not in archive.namelist():
                    warnings.append(
                        f"missing drawing part {drawing_part} for {title}:{drawing_rid}"
                    )
                    continue
                drawing_rels = _relationships(archive, drawing_part)
                drawing = ET.fromstring(archive.read(drawing_part))
                for xml_index, anchor in enumerate(list(drawing)):
                    anchor_kind = _local(anchor.tag)
                    if anchor_kind not in {"oneCellAnchor", "twoCellAnchor", "absoluteAnchor"}:
                        warnings.append(
                            f"unsupported anchor {anchor_kind} in {drawing_part} at {xml_index}"
                        )
                    anchor_from = _marker(anchor.find(f"{{{_XDR}}}from"))
                    anchor_to = _marker(anchor.find(f"{{{_XDR}}}to"))
                    pos = _marker(anchor.find(f"{{{_XDR}}}pos"))
                    extent = _marker(anchor.find(f"{{{_XDR}}}ext"))
                    if anchor_from is None:
                        anchor_from = pos
                    siblings = [
                        child
                        for child in anchor
                        if _local(child.tag) not in {"from", "to", "pos", "ext", "clientData"}
                    ]
                    if not siblings:
                        warnings.append(f"empty anchor in {drawing_part} at {xml_index}")
                    for sibling_index, child in enumerate(siblings):
                        objects.extend(
                            _walk_objects(
                                element=child,
                                archive=archive,
                                workbook_hash=payload_hash,
                                worksheet=title,
                                worksheet_part=worksheet_part,
                                drawing_part=drawing_part,
                                drawing_rid=drawing_rid,
                                drawing_rels=drawing_rels,
                                xml_index=xml_index,
                                z_index=xml_index,
                                anchor_kind=anchor_kind,
                                anchor_from=anchor_from,
                                anchor_to=anchor_to,
                                extent=extent,
                                edit_as=anchor.get("editAs"),
                                group_path=(sibling_index,),
                            )
                        )
    return DrawingManifest(payload_hash, tuple(objects), tuple(warnings))


def _bbox(item: DrawingObject) -> tuple[int, int, int, int] | None:
    start = item.anchor_from
    if start is None:
        return None
    if start.x is not None and start.y is not None:
        if item.extent and item.extent.cx is not None and item.extent.cy is not None:
            return start.x, start.y, start.x + item.extent.cx, start.y + item.extent.cy
        return None
    if start.col is None or start.row is None:
        return None
    # Used only by conservative grouping. Exact source markers remain in metadata.
    x0 = start.col * 609_600 + (start.col_off or 0)
    y0 = start.row * 190_500 + (start.row_off or 0)
    end = item.anchor_to
    if end and end.col is not None and end.row is not None:
        return (
            x0,
            y0,
            end.col * 609_600 + (end.col_off or 0),
            end.row * 190_500 + (end.row_off or 0),
        )
    if item.extent and item.extent.cx is not None and item.extent.cy is not None:
        return x0, y0, x0 + item.extent.cx, y0 + item.extent.cy
    return None


def _overlap_ratio(a: tuple[int, int, int, int], b: tuple[int, int, int, int]) -> float:
    intersection = max(0, min(a[2], b[2]) - max(a[0], b[0])) * max(
        0, min(a[3], b[3]) - max(a[1], b[1])
    )
    area = max(1, (a[2] - a[0]) * (a[3] - a[1]))
    return intersection / area


def _is_overlay(item: DrawingObject) -> bool:
    if item.kind == "cxnSp":
        return True
    if item.kind != "sp":
        return False
    geometry = item.properties.get("geometry")
    has_text = bool(item.text_runs)
    line = item.properties.get("line")
    has_visible_line = (
        isinstance(line, dict)
        and isinstance(line.get("fill"), dict)
        and line["fill"].get("kind") != "noFill"
    )
    fill = item.properties.get("fill")
    has_visible_fill = isinstance(fill, dict) and fill.get("kind") != "noFill"
    return (
        has_text
        or has_visible_line
        or has_visible_fill
        or geometry in {"line", "straightConnector1"}
    )


def group_drawing_components(
    manifest: DrawingManifest,
    *,
    blocked_object_ids: frozenset[str] = frozenset(),
) -> tuple[DrawingComponent, ...]:
    """Group only explicit groups or unambiguous overlays on one dominant picture."""
    by_anchor: dict[tuple[str, str, int], list[DrawingObject]] = {}
    roots = [item for item in manifest.objects if len(item.group_path) == 1]
    roots_by_part: dict[tuple[str, str], list[DrawingObject]] = {}
    bbox_by_id = {item.object_id: _bbox(item) for item in roots}
    for root in roots:
        roots_by_part.setdefault((root.worksheet, root.drawing_part), []).append(root)
    for item in manifest.objects:
        by_anchor.setdefault((item.worksheet, item.drawing_part, item.xml_index), []).append(item)
    consumed: set[str] = set()
    components: list[DrawingComponent] = []
    for root in roots:
        if root.object_id in consumed or root.object_id in blocked_object_ids:
            continue
        anchor_members = by_anchor[(root.worksheet, root.drawing_part, root.xml_index)]
        if root.kind == "grpSp":
            members = tuple(
                member for member in anchor_members if member.object_id not in blocked_object_ids
            )
            consumed.update(member.object_id for member in members)
            source_fidelity = (
                "exact_single_picture"
                if sum(member.is_picture for member in members) == 1 and len(members) == 2
                else "members_retained"
            )
            warning = (
                None
                if source_fidelity == "exact_single_picture"
                else "group composite requires calibrated worksheet rendering; members retained"
            )
            components.append(_component(manifest.workbook_hash, members, source_fidelity, warning))
            continue
        if root.is_picture:
            root_box = bbox_by_id[root.object_id]
            overlays: list[DrawingObject] = []
            part_roots = roots_by_part[(root.worksheet, root.drawing_part)]
            part_pictures = [item for item in part_roots if item.is_picture]
            if root_box is not None:
                candidates = [
                    item
                    for item in part_roots
                    if item.object_id not in blocked_object_ids
                    and item.object_id not in consumed
                    and item.object_id != root.object_id
                    and _is_overlay(item)
                    and bbox_by_id[item.object_id] is not None
                    and _overlap_ratio(bbox_by_id[item.object_id], root_box) >= 0.5  # type: ignore[arg-type]
                ]
                for overlay in candidates:
                    overlay_box = bbox_by_id[overlay.object_id]
                    dominant = [
                        picture
                        for picture in part_pictures
                        if bbox_by_id[picture.object_id] is not None
                        and overlay_box is not None
                        and _overlap_ratio(overlay_box, bbox_by_id[picture.object_id]) >= 0.5  # type: ignore[arg-type]
                    ]
                    if dominant == [root]:
                        overlays.append(overlay)
            members = tuple(
                sorted((root, *overlays), key=lambda item: (item.z_index, item.group_path))
            )
            consumed.update(member.object_id for member in members)
            fidelity = "exact_single_picture" if len(members) == 1 else "members_retained"
            warning = (
                None
                if len(members) == 1
                else (
                    "overlay composite requires calibrated worksheet rendering; "
                    "source members retained"
                )
            )
            components.append(_component(manifest.workbook_hash, members, fidelity, warning))
            continue
        consumed.add(root.object_id)
    return tuple(components)


def _component(
    workbook_hash: str,
    members: tuple[DrawingObject, ...],
    source_fidelity: str,
    warning: str | None,
) -> DrawingComponent:
    first = members[0]
    identity = ":".join(
        (workbook_hash, first.worksheet, first.drawing_part, *[item.object_id for item in members])
    )
    return DrawingComponent(
        component_id=hashlib.sha256(identity.encode()).hexdigest(),
        worksheet=first.worksheet,
        drawing_part=first.drawing_part,
        members=members,
        source_fidelity=source_fidelity,
        warning=warning,
    )


def anchor_cell_range(item: DrawingObject) -> str | None:
    start = item.anchor_from
    if start is None or start.col is None or start.row is None:
        return None
    end = item.anchor_to
    first = f"{get_column_letter(start.col + 1)}{start.row + 1}"
    if end is None or end.col is None or end.row is None:
        return first
    last = f"{get_column_letter(end.col + 1)}{end.row + 1}"
    return first if first == last else f"{first}:{last}"
